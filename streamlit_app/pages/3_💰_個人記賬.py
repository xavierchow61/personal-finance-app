"""個人記賬 — 分錄、帳戶總覽及手動分錄"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import streamlit as st
import pandas as pd

from streamlit_app._common import C, app_header, init_dbs, render_subpage_nav

st.set_page_config(
    page_title="個人記賬", page_icon="💰", layout="wide",
    initial_sidebar_state="expanded",
)
init_dbs()

app_header("個人記賬", "💰", "雙式記賬（Double-Entry Ledger）")
render_subpage_nav("ledger")

from personal_finance import db as pfdb, reports as pfr
from personal_finance import excel_export as pfexp
import cached_pfr

# === 🚀 並行 fetch（一次 spinner，所有資料一起載入）===
if True:   # ⚡ 移除 spinner（cache hit 即時，避免閃動）
    _bundle = cached_pfr.fetch_ledger_bundle()
_entries = _bundle["entries"]
_accounts = _bundle["accounts"]
_balances = _bundle["balances"]

# === 頂部工具列：匯出 Excel（一鍵下載）===
def _build_pf_xlsx_direct(period):
    """建立個人記賬 Excel；用 BytesIO（避免 /tmp 權限/位置問題）
    無 cache — 確保 error 即時可見
    """
    import io
    from openpyxl import Workbook
    from personal_finance.excel_export import (
        _write_overview, _write_accounts, _write_journal,
        _write_budget, _write_projects, _write_fx_rates,
        _write_aliases,
    )
    buf = io.BytesIO()
    wb = Workbook()
    if wb.active:
        wb.remove(wb.active)
    # 每個 _write_xxx 個別 try，唔好整個 fail
    sections = [
        ("overview", lambda: _write_overview(wb, period=period)),
        ("accounts", lambda: _write_accounts(wb, period=period)),
        ("journal", lambda: _write_journal(wb, period=period)),
        ("budget", lambda: _write_budget(wb, period=period)),
        ("projects", lambda: _write_projects(wb)),
        ("fx_rates", lambda: _write_fx_rates(wb)),
        ("aliases", lambda: _write_aliases(wb)),
    ]
    errors = []
    for name, fn in sections:
        try:
            fn()
        except Exception as ex:
            errors.append(f"{name}: {type(ex).__name__}: {ex}")
    if not wb.worksheets:
        # 完全冇任何 sheet → 加個 placeholder
        ws = wb.create_sheet("空白")
        ws["A1"] = "（無資料可匯出）"
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), errors


exp_l, exp_c, exp_r = st.columns([2, 3, 2])
with exp_c:
    period_in = st.text_input("📅 指定 period（選填，YYYY-MM）", "",
                                placeholder="留空 = 全部")
with exp_r:
    st.markdown("&nbsp;", unsafe_allow_html=True)  # 對齊 label 高度
    _pf_xlsx_key = f"_pf_xlsx_{period_in or 'all'}"
    if st.session_state.get(_pf_xlsx_key):
        import datetime as _dt
        ts = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        st.download_button(
            "⬇️ 下載 Excel",
            data=st.session_state[_pf_xlsx_key],
            file_name=f"個人記賬_{period_in or '全部'}_{ts}.xlsx",
            mime=("application/vnd.openxmlformats-officedocument."
                  "spreadsheetml.sheet"),
            use_container_width=True,
            type="primary",
            key=f"dl_pf_{_pf_xlsx_key}",
        )
    else:
        if st.button("📊 準備匯出 Excel",
                      use_container_width=True,
                      type="primary",
                      key="prep_pf_xlsx"):
            st.toast("📊 開始建立 Excel...", icon="⏳")
            try:
                _pf_data, _section_errors = _build_pf_xlsx_direct(
                    period_in.strip() or None)
                if not _pf_data:
                    st.error("Excel build 完成但 byte 為空")
                else:
                    st.session_state[_pf_xlsx_key] = _pf_data
                    if _section_errors:
                        st.warning(
                            "⚠️ 部分 sheet 建立失敗（但其他 OK）：\n\n"
                            + "\n\n".join(_section_errors[:3])
                        )
                    st.toast(
                        f"✅ Excel 已建立 ({len(_pf_data)//1024} KB)",
                        icon="✅",
                    )
                    st.rerun()
            except Exception as ex:
                import traceback
                st.error(
                    f"❌ 匯出失敗：{type(ex).__name__}: {ex}"
                )
                with st.expander("🔍 Stack trace（debug 用）"):
                    st.code(traceback.format_exc())

st.divider()

tab1, tab2, tab3 = st.tabs(["📜 交易紀錄", "🏦 帳戶總覽", "➕ 新增分錄"])

# === 交易紀錄 ===
with tab1:
    st.caption("最新 100 筆分錄")
    entries = _entries   # 用 bundle 資料
    if entries:
        df = pd.DataFrame([
            {
                "ID": e["entry_id"],
                "日期": e["entry_date"],
                "說明": e["description"] or "",
                "幣別": e.get("currency") or "HKD",
                "匯率": e.get("fx_rate") or 1.0,
                "金額（原幣）": e.get("amount", 0),
                "分錄細項": e.get("lines_summary") or "",
            }
            for e in entries
        ])
        st.caption("💡 點擊任何一行查看詳細分錄")
        sel_je = st.dataframe(
            df, hide_index=True, use_container_width=True,
            on_select="rerun", selection_mode="single-row",
            column_config={
                "金額（原幣）": st.column_config.NumberColumn(
                    format="$%.2f"),
                "匯率": st.column_config.NumberColumn(format="%.4f"),
            },
            key="je_table",
        )

        # === 選定行 → 顯示詳細 ===
        if sel_je.selection.rows:
            sel_je_id = int(df.iloc[sel_je.selection.rows[0]]["ID"])
            full = pfdb.get_entry(sel_je_id)
            if full:
                st.divider()
                d1, d2 = st.columns([3, 1])
                with d1:
                    st.markdown(
                        f"### 📋 分錄 #{sel_je_id}：{full.get('description', '')}"
                    )
                    st.caption(
                        f"📅 {full['entry_date']}"
                        + (f" · 🧾 關聯單據 #{full['invoice_id']}"
                           if full.get('invoice_id') else "")
                    )
                with d2:
                    if st.button("🗑️ 刪除此分錄",
                                  type="secondary",
                                  use_container_width=True,
                                  key=f"del_je_{sel_je_id}"):
                        try:
                            pfdb.delete_entry(sel_je_id)
                            st.toast(
                                f"🗑️ 已刪除分錄 #{sel_je_id}",
                                icon="🗑️")
                            try:
                                import cached_pfr
                                cached_pfr.invalidate_journal()
                            except Exception:
                                pass
                            st.rerun()
                        except Exception as ex:
                            st.error(str(ex))
                line_df = pd.DataFrame([
                    {
                        "帳戶": l["account_code"],
                        "借方": l["debit"] or "",
                        "貸方": l["credit"] or "",
                    }
                    for l in full["lines"]
                ])
                st.dataframe(line_df, hide_index=True,
                              use_container_width=True)
    else:
        st.info("尚無分錄。請先前往「📤 提取單據」或於下方手動新增。")

# === 帳戶總覽 ===
with tab2:
    st.caption("所有帳戶及目前餘額（以港幣顯示）— 父帳戶顯示本身與子帳戶合計")
    from streamlit_app._common import group_accounts_with_parent_totals
    # 結合 balance 入 _accounts
    bal_map = {b["code"]: b["balance"] for b in _balances}
    accs_with_bal = [
        {**a, "balance": bal_map.get(a["code"], 0)}
        for a in _accounts
        if a.get("is_active")  # 只顯示啟用嘅
    ]
    grouped = group_accounts_with_parent_totals(accs_with_bal)

    df_rows = []
    for item in grouped:
        a = item["account"]
        depth = item["depth"]
        is_parent = item["is_parent"]
        is_indent = depth > 0
        name_prefix = "　└ " if is_indent else ""
        balance_label = (
            "本身餘額" if is_parent else "目前餘額（HKD）"
        )
        # 父帳戶（有子）顯示合計；其他顯示自己餘額
        display_balance = (
            item["aggregated_balance"] if is_parent
            else float(a.get("balance") or 0)
        )
        row = {
            "代碼": a["code"],
            "名稱": (
                f"{name_prefix}{a.get('icon') or ''} {a['name']}"
                + (f"（合計 {item['n_children']} 子）"
                   if is_parent else "")
            ),
            "類型": a["account_type"],
            "幣別": a.get("currency") or "HKD",
            "期初餘額": a.get("opening_balance") or 0,
            "目前餘額（HKD）": display_balance,
            "啟用中": "✅" if a["is_active"] else "❌",
        }
        df_rows.append(row)

    df = pd.DataFrame(df_rows)
    st.dataframe(df, hide_index=True, use_container_width=True,
                  column_config={
                      "期初餘額": st.column_config.NumberColumn(format="$%.2f"),
                      "目前餘額（HKD）": st.column_config.NumberColumn(
                          format="$%.2f"),
                  })

# === 新增手動分錄 ===
with tab3:
    st.subheader("➕ 新增手動分錄")
    st.caption("例如：薪金、帳戶轉移、信用卡還款")

    from streamlit_app._common import cascade_account_picker

    from datetime import date as _d
    me1, me2 = st.columns(2)
    with me1:
        edate = st.date_input("日期", _d.today())
        amount = st.number_input("金額", value=0.0, format="%.2f")
        st.markdown("**借方帳戶（Dr）**")
        from_acc_code = cascade_account_picker(
            "借方", key_prefix="entry_dr",
        )
    with me2:
        desc = st.text_input("說明", "")
        currency = st.selectbox("幣別",
                                 ["HKD", "USD", "JPY", "CNY", "EUR", "GBP"],
                                 index=0)
        st.markdown("**貸方帳戶（Cr）**")
        to_acc_code = cascade_account_picker(
            "貸方", key_prefix="entry_cr",
        )

    if st.button("💾 寫入分錄", type="primary"):
        if amount <= 0:
            st.error("金額必須大於零")
        elif not from_acc_code or not to_acc_code:
            st.error("請選擇借方與貸方帳戶")
        elif from_acc_code == to_acc_code:
            st.error("借方與貸方不能是同一帳戶")
        else:
            try:
                eid = pfdb.create_entry(
                    entry_date=edate.isoformat(),
                    description=desc or
                        f"手動分錄：{from_acc_code} → {to_acc_code}",
                    lines=[
                        {"account_code": from_acc_code,
                         "debit": amount, "credit": 0},
                        {"account_code": to_acc_code,
                         "debit": 0, "credit": amount},
                    ],
                    currency=currency,
                )
                st.toast(f"✅ 已寫入分錄 #{eid}", icon="✅")
                try:
                    import cached_pfr
                    cached_pfr.invalidate_journal()
                except Exception:
                    pass
                st.rerun()
            except Exception as ex:
                st.error(str(ex))
