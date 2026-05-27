"""單據紀錄 — 列出全部已提取的單據，可篩選及編輯"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import streamlit as st
import pandas as pd

from streamlit_app._common import C, app_header, init_dbs, render_subpage_nav

st.set_page_config(
    page_title="單據紀錄", page_icon="📋", layout="wide",
    initial_sidebar_state="expanded",
)
init_dbs()

app_header("單據處理", "📋", "所有經 AI 提取的單據資料")
render_subpage_nav("invoice")

import database as invdb
import excel_exporter
from config import CATEGORIES

# === 頂部工具列：匯出 Excel ===
top_l, top_r = st.columns([5, 2])
with top_r:
    if st.button("📊 匯出全部單據為 Excel", use_container_width=True):
        try:
            import io, datetime as _dt
            from openpyxl import Workbook
            buf = io.BytesIO()
            wb = Workbook()
            dash = wb.active
            invoices = invdb.list_all()
            excel_exporter._write_dashboard(dash, invoices)
            main = wb.create_sheet()
            excel_exporter._write_main_sheet(main, invoices)
            reimb = wb.create_sheet()
            excel_exporter._write_reimbursement_sheet(reimb)
            wb.active = 0
            wb.save(buf)
            buf.seek(0)
            st.session_state["xlsx_buffer"] = buf.getvalue()
            st.session_state["xlsx_name"] = (
                f"單據紀錄_{_dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
            )
        except Exception as ex:
            st.error(f"匯出失敗：{ex}")

    if "xlsx_buffer" in st.session_state:
        st.download_button(
            "⬇️ 點此下載 Excel",
            data=st.session_state["xlsx_buffer"],
            file_name=st.session_state.get("xlsx_name",
                                              "單據紀錄.xlsx"),
            mime=("application/vnd.openxmlformats-officedocument."
                  "spreadsheetml.sheet"),
            use_container_width=True,
        )

# === 篩選器 ===
fc1, fc2, fc3, fc4 = st.columns(4)
search = fc1.text_input("🔍 搜尋商戶 / 備註")
cat_filter = fc2.selectbox("類別", ["（全部）"] + CATEGORIES)
exp_filter = fc3.selectbox("類型", ["（全部）", "私人", "公司報銷", "可扣稅"])
reimb_filter = fc4.selectbox("報銷狀態",
                              ["（全部）", "✅ 已收款", "⏳ 待收款"])

# === 載入並篩選 ===
all_invoices = invdb.list_all()
filtered = all_invoices
if search:
    s = search.lower()
    filtered = [i for i in filtered
                 if s in (i.get("store_name") or "").lower()
                 or s in (i.get("notes") or "").lower()]
if cat_filter != "（全部）":
    filtered = [i for i in filtered if i.get("category") == cat_filter]
if exp_filter != "（全部）":
    filtered = [i for i in filtered
                 if (i.get("expense_type") or "私人") == exp_filter]
if reimb_filter == "✅ 已收款":
    filtered = [i for i in filtered if i.get("reimbursed")]
elif reimb_filter == "⏳ 待收款":
    filtered = [i for i in filtered if not i.get("reimbursed")
                 and i.get("expense_type") == "公司報銷"]

st.caption(f"共 {len(filtered)} 張單據（總數 {len(all_invoices)} 張）")

# === 建立資料表 ===
if filtered:
    EXPENSE_ICONS = {"私人": "👤", "公司報銷": "🏢", "可扣稅": "🧾"}
    df = pd.DataFrame([
        {
            "ID": inv["id"],
            "類型": EXPENSE_ICONS.get(inv.get("expense_type") or "私人", "👤"),
            "日期": inv.get("purchase_date") or "—",
            "商戶": inv.get("store_name") or "—",
            "類別": inv.get("category") or "—",
            "金額": inv.get("total_amount") or 0,
            "幣別": inv.get("currency") or "—",
            "付款方式": inv.get("payment_method") or "—",
            "已收款": ("✅" if inv.get("reimbursed")
                        else ("⏳" if inv.get("expense_type") == "公司報銷"
                              else "—")),
        }
        for inv in filtered
    ])
    selected = st.dataframe(
        df, hide_index=True, use_container_width=True,
        column_config={
            "金額": st.column_config.NumberColumn(format="$%.2f"),
        },
        on_select="rerun", selection_mode="single-row",
    )

    # 如果選取了一行
    if selected.selection.rows:
        idx = selected.selection.rows[0]
        inv_id = int(df.iloc[idx]["ID"])
        inv = invdb.get_invoice(inv_id)

        st.divider()

        # === 兩欄佈局：左邊編輯表單，右邊單據圖預覽 ===
        edit_col, preview_col = st.columns([3, 2])

        with preview_col:
            st.markdown(f"### 📷 單據預覽 #{inv['id']}")
            src = inv.get("source_file")
            src_path = Path(src) if src else None
            image_shown = False

            if src_path and src_path.exists():
                suffix = src_path.suffix.lower()
                if suffix in {".jpg", ".jpeg", ".png", ".bmp",
                              ".webp", ".tiff"}:
                    st.image(str(src_path),
                              caption=src_path.name,
                              use_container_width=True)
                    image_shown = True
                elif suffix == ".pdf":
                    # 嘗試 render PDF 第一頁
                    try:
                        from pdf_utils import pdf_first_page_to_image
                        img = pdf_first_page_to_image(src_path)
                        st.image(img,
                                  caption=f"📄 {src_path.name}（第 1 頁）",
                                  use_container_width=True)
                        image_shown = True
                    except Exception as ex:
                        st.info(
                            f"📄 PDF：{src_path.name}\n\n"
                            f"無法 render 預覽：{ex}"
                        )

            # 圖片無法顯示 → 允許重新上傳
            if not image_shown:
                if src and not (src_path and src_path.exists()):
                    st.warning(
                        f"⚠️ 原始檔案已不存在或位置改變"
                    )
                    st.caption(f"原路徑：`{src}`")
                else:
                    st.caption("（此單據沒有關聯的原始檔案）")

                # 重新上傳介面
                new_img = st.file_uploader(
                    "📷 重新上傳此單據圖片",
                    type=["jpg", "jpeg", "png", "bmp",
                          "webp", "tiff", "pdf"],
                    key=f"reup_{inv_id}",
                )
                if new_img is not None:
                    # 儲存到 outputs/invoice_images/
                    save_dir = (Path(__file__).resolve().parents[2]
                                / "outputs" / "invoice_images")
                    save_dir.mkdir(parents=True, exist_ok=True)
                    save_path = save_dir / (
                        f"{inv_id}_{new_img.name}"
                    )
                    save_path.write_bytes(new_img.getvalue())
                    # 更新 DB
                    updated = dict(inv)
                    updated["source_file"] = str(save_path)
                    invdb.update_invoice(inv_id, updated)
                    st.success(
                        f"✅ 已關聯新圖片：{save_path.name}"
                    )
                    st.rerun()

        with edit_col:
            st.subheader(f"✏️ 編輯 #{inv['id']}")
            ec1, ec2 = st.columns(2)
            with ec1:
                store = st.text_input("商戶",
                                       inv.get("store_name") or "")
                date = st.text_input("日期（YYYY-MM-DD）",
                                      inv.get("purchase_date") or "")
                amount = st.number_input(
                    "金額",
                    value=float(inv.get("total_amount") or 0),
                    format="%.2f")
                cat = st.selectbox(
                    "類別", CATEGORIES,
                    index=(CATEGORIES.index(inv.get("category"))
                           if inv.get("category") in CATEGORIES else 0))
            with ec2:
                currency = st.text_input("幣別",
                                          inv.get("currency") or "HKD")

                # === 付款方式：dropdown + 自訂選項 ===
                from personal_finance import db as _pfdb
                alias_keywords = sorted({
                    a["keyword"] for a in _pfdb.list_payment_aliases()
                })
                past_methods = sorted({
                    i.get("payment_method") or ""
                    for i in all_invoices
                    if i.get("payment_method")
                })
                payment_options = sorted(
                    set(alias_keywords) | set(past_methods)
                )
                # 確保目前值在 options 中
                cur_payment = inv.get("payment_method") or ""
                if cur_payment and cur_payment not in payment_options:
                    payment_options.insert(0, cur_payment)
                payment_options = ["✏️ 自訂輸入..."] + payment_options

                # 找出目前值的 index
                if cur_payment in payment_options:
                    pay_idx = payment_options.index(cur_payment)
                else:
                    pay_idx = 0
                payment_sel = st.selectbox(
                    "付款方式",
                    payment_options,
                    index=pay_idx,
                    help=f"已知選項：{len(payment_options)-1} 個"
                          f"（含付款對應 + 過往單據）",
                )
                if payment_sel == "✏️ 自訂輸入...":
                    payment = st.text_input(
                        "輸入新的付款方式",
                        value=cur_payment,
                        key=f"custom_pay_{inv_id}",
                    )
                else:
                    payment = payment_sel
                exp_type = st.selectbox(
                    "類型", ["私人", "公司報銷", "可扣稅"],
                    index=["私人", "公司報銷", "可扣稅"].index(
                        inv.get("expense_type") or "私人"))
                reimb = st.checkbox("已收款",
                                     value=bool(inv.get("reimbursed")))
            notes = st.text_area("備註", inv.get("notes") or "")

            # === AR Receive Dialog（勾「已收款」且係公司報銷時顯示）===
            ar_acc_code = None
            ar_date_iso = None
            was_reimbursed = bool(inv.get("reimbursed"))
            if (exp_type == "公司報銷" and reimb and not was_reimbursed):
                st.info(
                    "💰 偵測到你勾選「已收款」"
                    "— 請選擇收款入哪個帳戶，"
                    "系統會自動入收款分錄。"
                )
                from personal_finance import db as pfdb
                from datetime import date as _d
                asset_accs = pfdb.list_accounts(account_type="asset")
                ar_opts = {
                    f"{a.get('icon') or ''} {a['name']} ({a['code']})":
                        a["code"]
                    for a in asset_accs
                }
                ar_a, ar_b = st.columns(2)
                with ar_a:
                    ar_acc_label = st.selectbox(
                        "💳 收款帳戶", list(ar_opts.keys()),
                        key=f"ar_dlg_acc_{inv_id}")
                    ar_acc_code = ar_opts[ar_acc_label]
                with ar_b:
                    ar_date = st.date_input(
                        "📅 收款日期", _d.today(),
                        key=f"ar_dlg_date_{inv_id}")
                    ar_date_iso = ar_date.isoformat()

            save_col, del_col = st.columns(2)
            if save_col.button("💾 儲存修改", type="primary",
                                use_container_width=True):
                old_type = (inv.get("expense_type") or "私人").strip()
                old_amount = float(inv.get("total_amount") or 0)
                old_cat = inv.get("category") or ""
                old_payment = inv.get("payment_method") or ""

                data = dict(inv)
                data.update({
                    "store_name": store, "purchase_date": date,
                    "total_amount": amount, "category": cat,
                    "currency": currency, "payment_method": payment,
                    "expense_type": exp_type, "reimbursed": reimb,
                    "notes": notes,
                })
                invdb.update_invoice(inv_id, data)

                needs_repost = (
                    exp_type != old_type
                    or abs(amount - old_amount) > 0.005
                    or cat != old_cat
                    or payment != old_payment
                )

                posting_msg = ""
                if needs_repost:
                    try:
                        from personal_finance import posting as pfpost
                        n_removed = pfpost.unpost_invoice(inv_id)
                        new_entry_id = pfpost.post_invoice(data)
                        if exp_type == "公司報銷":
                            posting_msg = (
                                f" · 🏢 已自動入 AR_REIMBURSE "
                                f"（刪 {n_removed} 條舊分錄、"
                                f"新分錄 #{new_entry_id}）"
                            )
                        else:
                            posting_msg = (
                                f" · 💰 已重新入賬"
                                f"（刪 {n_removed} 條舊分錄、"
                                f"新分錄 #{new_entry_id}）"
                            )
                    except Exception as ex:
                        st.error(f"⚠️ 重新入賬失敗：{ex}")

                # === AR 收款入賬（新勾「已收款」+ 公司報銷時）===
                if (exp_type == "公司報銷" and reimb
                        and not was_reimbursed and ar_acc_code):
                    try:
                        from personal_finance import posting as pfpost
                        rid = pfpost.mark_reimbursement_received(
                            invoice_id=inv_id,
                            received_account=ar_acc_code,
                            received_date=ar_date_iso,
                        )
                        posting_msg += (
                            f" · ✅ 收款入賬"
                            f"（Dr {ar_acc_code} / Cr AR_REIMBURSE，"
                            f"分錄 #{rid}）"
                        )
                    except Exception as ex:
                        st.error(f"⚠️ AR 收款入賬失敗：{ex}")

                st.success(f"✅ 已更新 #{inv_id}{posting_msg}")
                st.rerun()
            if del_col.button("🗑️ 刪除", type="secondary",
                               use_container_width=True):
                invdb.delete_invoice(inv_id)
                st.success(f"已刪除 #{inv_id}")
                st.rerun()
else:
    st.info("並無符合條件的單據。請先前往「📤 提取單據」匯入資料。")
