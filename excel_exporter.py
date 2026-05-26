"""匯出 Excel：明細表 + Dashboard + 報銷單"""
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import database
from config import OUTPUT_DIR


# === 樣式 ===
HEADER_FILL = PatternFill(start_color="1E66F5", end_color="1E66F5", fill_type="solid")
REIMB_HEADER_FILL = PatternFill(start_color="287D22", end_color="287D22", fill_type="solid")
HEADER_FONT = Font(name="Microsoft JhengHei UI", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Microsoft JhengHei UI", size=16, bold=True, color="1E66F5")
SUBTITLE_FONT = Font(name="Microsoft JhengHei UI", size=11, bold=True, color="287D22")
WARN_FONT = Font(name="Microsoft JhengHei UI", size=11, bold=True, color="BC7700")
NORMAL_FONT = Font(name="Microsoft JhengHei UI", size=10)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)
ALT_FILL = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
PENDING_FILL = PatternFill(start_color="FFF4E0", end_color="FFF4E0", fill_type="solid")


HEADERS = [
    ("id", "ID", 6),
    ("purchase_date", "購買日期", 13),
    ("store_name", "商店名稱", 22),
    ("category", "類別", 12),
    ("expense_type", "支出類型", 12),
    ("reimbursed", "已報銷", 10),
    ("total_amount", "總金額", 12),
    ("currency", "幣值", 8),
    ("payment_method", "付款方式", 14),
    ("items_summary", "產品內容", 40),
    ("tax", "稅項", 8),
    ("receipt_number", "單號", 18),
    ("notes", "備註", 25),
    ("source_file", "原檔案", 35),
]


def _items_summary(items: list[dict]) -> str:
    if not items:
        return ""
    parts = []
    for it in items[:10]:
        name = it.get("name", "")
        qty = it.get("quantity", "")
        price = it.get("price", "")
        bits = [str(name)]
        if qty:
            bits.append(f"x{qty}")
        if price:
            bits.append(f"${price}")
        parts.append(" ".join(bits))
    if len(items) > 10:
        parts.append(f"...(+{len(items)-10})")
    return "; ".join(parts)


def _reimb_str(inv: dict) -> str:
    etype = inv.get("expense_type") or "私人"
    if etype != "公司報銷":
        return "—"
    return "✅ 已報銷" if inv.get("reimbursed") else "⏳ 待報銷"


def _write_main_sheet(ws, invoices: list[dict]):
    ws.title = "📋 單據明細"

    for col_idx, (_, label, width) in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 26

    for r, inv in enumerate(invoices, 2):
        values = {
            "id": inv.get("id"),
            "purchase_date": inv.get("purchase_date"),
            "store_name": inv.get("store_name"),
            "category": inv.get("category"),
            "expense_type": inv.get("expense_type") or "私人",
            "reimbursed": _reimb_str(inv),
            "total_amount": inv.get("total_amount"),
            "currency": inv.get("currency"),
            "payment_method": inv.get("payment_method"),
            "items_summary": _items_summary(inv.get("items") or []),
            "tax": inv.get("tax"),
            "receipt_number": inv.get("receipt_number"),
            "notes": inv.get("notes"),
            "source_file": inv.get("source_file"),
        }
        for col_idx, (key, _, _) in enumerate(HEADERS, 1):
            cell = ws.cell(row=r, column=col_idx, value=values[key])
            cell.font = NORMAL_FONT
            cell.border = THIN_BORDER
            if key in ("total_amount", "tax"):
                cell.alignment = RIGHT
                cell.number_format = "#,##0.00"
            elif key in ("id", "currency", "purchase_date", "category",
                          "expense_type", "reimbursed"):
                cell.alignment = CENTER
            else:
                cell.alignment = LEFT
            if r % 2 == 0:
                cell.fill = ALT_FILL
            # 待報銷標黃
            if (key == "reimbursed" and values["reimbursed"] == "⏳ 待報銷"):
                cell.fill = PENDING_FILL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_dashboard(ws, invoices: list[dict]):
    ws.title = "📊 Dashboard"

    ws["A1"] = "📊 支出 Dashboard"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:H1")
    ws["A1"].alignment = LEFT

    ws["A2"] = f"生成時間：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(name="Microsoft JhengHei UI", size=9, color="888888")
    ws.merge_cells("A2:H2")

    # === KPI 區 ===
    total = database.grand_total()
    cnt = database.count()
    avg = (total / cnt) if cnt else 0
    reimb = database.reimbursement_summary()

    kpi_row = 4
    kpis = [
        ("總單據數", cnt, "#,##0", "1E66F5"),
        ("總支出", total, "#,##0.00", "1E66F5"),
        ("平均每單", avg, "#,##0.00", "1E66F5"),
        ("🏢 未報銷", reimb["company_pending"], "#,##0.00", "BC7700"),
        ("🧾 可扣稅", reimb["tax_deductible_total"], "#,##0.00", "287D22"),
    ]
    for i, (label, value, fmt, color) in enumerate(kpis):
        col = 1 + i * 2
        lbl_cell = ws.cell(row=kpi_row, column=col, value=label)
        lbl_cell.font = Font(name="Microsoft JhengHei UI", size=10, color="5C5F77")
        lbl_cell.alignment = CENTER
        val_cell = ws.cell(row=kpi_row + 1, column=col, value=value)
        val_cell.font = Font(name="Microsoft JhengHei UI", size=16, bold=True, color=color)
        val_cell.alignment = CENTER
        val_cell.number_format = fmt
        ws.merge_cells(start_row=kpi_row, start_column=col,
                       end_row=kpi_row, end_column=col + 1)
        ws.merge_cells(start_row=kpi_row + 1, start_column=col,
                       end_row=kpi_row + 1, end_column=col + 1)
    ws.row_dimensions[kpi_row + 1].height = 30

    # === Section 1: Top 5 最貴單據 ===
    sec_row = 8
    ws.cell(row=sec_row, column=1, value="🏆 Top 5 最貴單據").font = SUBTITLE_FONT
    ws.merge_cells(start_row=sec_row, start_column=1, end_row=sec_row, end_column=6)

    top_headers = ["排名", "購買日期", "商店", "類別", "金額", "幣值"]
    for c, h in enumerate(top_headers, 1):
        cell = ws.cell(row=sec_row + 1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    top5 = database.top_n_by_amount(5)
    for i, inv in enumerate(top5, 1):
        row = sec_row + 1 + i
        cells = [
            (1, i, CENTER, None),
            (2, inv.get("purchase_date"), CENTER, None),
            (3, inv.get("store_name") or "", LEFT, None),
            (4, inv.get("category") or "", CENTER, None),
            (5, inv.get("total_amount") or 0, RIGHT, "#,##0.00"),
            (6, inv.get("currency") or "", CENTER, None),
        ]
        for col, val, align, fmt in cells:
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = NORMAL_FONT
            cell.alignment = align
            cell.border = THIN_BORDER
            if fmt:
                cell.number_format = fmt
            if i % 2 == 0:
                cell.fill = ALT_FILL

    # === Section 2: Top 5 類別支出 ===
    cat_row = sec_row + 1 + max(len(top5), 1) + 3
    ws.cell(row=cat_row, column=1, value="💰 Top 5 類別支出").font = SUBTITLE_FONT
    ws.merge_cells(start_row=cat_row, start_column=1, end_row=cat_row, end_column=6)

    cat_headers = ["排名", "類別", "總支出", "佔比"]
    for c, h in enumerate(cat_headers, 1):
        cell = ws.cell(row=cat_row + 1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    top_cats = database.top_n_categories(5)
    total_for_pct = sum(t for _, t in database.category_totals()) or 1
    for i, (cat, amt) in enumerate(top_cats, 1):
        row = cat_row + 1 + i
        pct = amt / total_for_pct
        cells = [
            (1, i, CENTER, None),
            (2, cat, LEFT, None),
            (3, amt, RIGHT, "#,##0.00"),
            (4, pct, RIGHT, "0.00%"),
        ]
        for col, val, align, fmt in cells:
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = NORMAL_FONT
            cell.alignment = align
            cell.border = THIN_BORDER
            if fmt:
                cell.number_format = fmt
            if i % 2 == 0:
                cell.fill = ALT_FILL

    # === Section 3: 各類別佔比（圓餅圖 data table）===
    pie_data_row = cat_row + 1 + max(len(top_cats), 1) + 3
    ws.cell(row=pie_data_row, column=1, value="🥧 各類別佔比（圓餅圖）").font = SUBTITLE_FONT
    ws.merge_cells(start_row=pie_data_row, start_column=1, end_row=pie_data_row, end_column=6)

    pie_header_row = pie_data_row + 1
    ws.cell(row=pie_header_row, column=1, value="類別").fill = HEADER_FILL
    ws.cell(row=pie_header_row, column=2, value="總支出").fill = HEADER_FILL
    ws.cell(row=pie_header_row, column=1).font = HEADER_FONT
    ws.cell(row=pie_header_row, column=2).font = HEADER_FONT
    ws.cell(row=pie_header_row, column=1).alignment = CENTER
    ws.cell(row=pie_header_row, column=2).alignment = CENTER
    ws.cell(row=pie_header_row, column=1).border = THIN_BORDER
    ws.cell(row=pie_header_row, column=2).border = THIN_BORDER

    all_cats = database.category_totals()
    for i, (cat, amt) in enumerate(all_cats, 1):
        row = pie_header_row + i
        c1 = ws.cell(row=row, column=1, value=cat)
        c2 = ws.cell(row=row, column=2, value=amt)
        c1.font = NORMAL_FONT
        c2.font = NORMAL_FONT
        c1.alignment = LEFT
        c2.alignment = RIGHT
        c2.number_format = "#,##0.00"
        c1.border = THIN_BORDER
        c2.border = THIN_BORDER

    if all_cats:
        pie = PieChart()
        pie.title = "各類別支出佔比"
        labels = Reference(ws, min_col=1, min_row=pie_header_row + 1,
                           max_row=pie_header_row + len(all_cats))
        data = Reference(ws, min_col=2, min_row=pie_header_row,
                         max_row=pie_header_row + len(all_cats))
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.height = 12
        pie.width = 18
        pie.dataLabels = DataLabelList(showPercent=True, showCatName=True)
        ws.add_chart(pie, f"D{pie_data_row + 1}")

    if top_cats:
        bar = BarChart()
        bar.type = "bar"
        bar.style = 11
        bar.title = "Top 5 類別支出"
        bar.y_axis.title = "類別"
        bar.x_axis.title = "金額"
        bar.legend = None

        bar_labels = Reference(ws, min_col=2, min_row=cat_row + 2,
                               max_row=cat_row + 1 + len(top_cats))
        bar_data = Reference(ws, min_col=3, min_row=cat_row + 1,
                             max_row=cat_row + 1 + len(top_cats))
        bar.add_data(bar_data, titles_from_data=True)
        bar.set_categories(bar_labels)
        bar.height = 10
        bar.width = 18
        ws.add_chart(bar, f"H{cat_row + 1}")

    widths = [10, 22, 16, 14, 14, 10, 18, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_reimbursement_sheet(ws):
    """報銷單 sheet - 按月分組公司報銷單據"""
    ws.title = "🏢 公司報銷單"

    # 標題
    ws["A1"] = "🏢 公司報銷單"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:G1")
    ws["A1"].alignment = LEFT

    summary = database.reimbursement_summary()
    ws["A2"] = (f"待報銷：${summary['company_pending']:,.2f}  |  "
                f"已報銷：${summary['company_reimbursed']:,.2f}  |  "
                f"可扣稅：${summary['tax_deductible_total']:,.2f}")
    ws["A2"].font = WARN_FONT
    ws.merge_cells("A2:G2")

    company_inv = database.by_expense_type("公司報銷")
    tax_inv = database.by_expense_type("可扣稅")
    all_reimb = company_inv + tax_inv

    if not all_reimb:
        ws["A4"] = "（暫無公司報銷 / 可扣稅單據）"
        ws["A4"].font = NORMAL_FONT
        return

    # 按月分組
    by_month = defaultdict(list)
    for inv in all_reimb:
        date = inv.get("purchase_date") or "(無日期)"
        month = date[:7] if len(date) >= 7 else "(無日期)"
        by_month[month].append(inv)

    headers = ["ID", "日期", "商店", "類別", "金額", "幣值", "支出類型", "已報銷", "備註"]
    widths_r = [6, 12, 22, 12, 12, 8, 12, 12, 25]

    row = 4
    for month in sorted(by_month.keys(), reverse=True):
        month_invoices = by_month[month]
        month_total = sum(i.get("total_amount") or 0 for i in month_invoices)
        month_pending = sum(
            i.get("total_amount") or 0
            for i in month_invoices
            if i.get("expense_type") == "公司報銷" and not i.get("reimbursed")
        )

        # 月份小標題
        title_cell = ws.cell(row=row, column=1,
                              value=f"📅 {month}　總額 ${month_total:,.2f}　"
                                    f"待報銷 ${month_pending:,.2f}")
        title_cell.font = SUBTITLE_FONT
        title_cell.fill = PatternFill(start_color="E8F2D8", end_color="E8F2D8", fill_type="solid")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        row += 1

        # Header
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.fill = REIMB_HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = THIN_BORDER
        row += 1

        # Rows
        for i, inv in enumerate(month_invoices):
            etype = inv.get("expense_type") or "私人"
            data_row = [
                (1, inv.get("id"), CENTER, None),
                (2, inv.get("purchase_date") or "", CENTER, None),
                (3, inv.get("store_name") or "", LEFT, None),
                (4, inv.get("category") or "", CENTER, None),
                (5, inv.get("total_amount") or 0, RIGHT, "#,##0.00"),
                (6, inv.get("currency") or "", CENTER, None),
                (7, etype, CENTER, None),
                (8, _reimb_str(inv), CENTER, None),
                (9, inv.get("notes") or "", LEFT, None),
            ]
            for col, val, align, fmt in data_row:
                cell = ws.cell(row=row, column=col, value=val)
                cell.font = NORMAL_FONT
                cell.alignment = align
                cell.border = THIN_BORDER
                if fmt:
                    cell.number_format = fmt
                if i % 2 == 0:
                    cell.fill = ALT_FILL
                # 待報銷標黃
                if (col == 8 and etype == "公司報銷" and not inv.get("reimbursed")):
                    cell.fill = PENDING_FILL
            row += 1

        # Subtotal
        subtotal_cell = ws.cell(row=row, column=4, value="小計：")
        subtotal_cell.font = Font(name="Microsoft JhengHei UI", size=10, bold=True)
        subtotal_cell.alignment = RIGHT
        val_cell = ws.cell(row=row, column=5, value=month_total)
        val_cell.font = Font(name="Microsoft JhengHei UI", size=10, bold=True, color="287D22")
        val_cell.alignment = RIGHT
        val_cell.number_format = "#,##0.00"
        row += 2

    # 欄寬
    for i, w in enumerate(widths_r, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A4"


def export_excel(out_path: Path | str | None = None) -> Path:
    """匯出 Excel 包含 Dashboard + 明細 + 報銷單"""
    invoices = database.list_all()

    if out_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = OUTPUT_DIR / f"單據紀錄_{timestamp}.xlsx"
    else:
        out_path = Path(out_path)

    wb = Workbook()

    dash = wb.active
    _write_dashboard(dash, invoices)

    main = wb.create_sheet()
    _write_main_sheet(main, invoices)

    reimb = wb.create_sheet()
    _write_reimbursement_sheet(reimb)

    wb.active = 0

    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    path = export_excel()
    print(f"✅ 已匯出：{path}")
